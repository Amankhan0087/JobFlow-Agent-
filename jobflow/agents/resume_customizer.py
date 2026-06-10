"""
JobFlow Resume Customizer Agent
Customizes base resume for specific jobs using LLM, generates PDF and DOCX outputs.
"""

import json
import logging
import os
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from dotenv import load_dotenv

load_dotenv()

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from jobflow.database.db import fetch_one, fetch_all, insert_row, update_row, init_db

logger = logging.getLogger(__name__)

BASE_RESUME_PATH = Path(__file__).parent.parent / "resume" / "base_resume.json"
RESUME_OUTPUT_DIR = Path(__file__).parent.parent.parent / "resume_outputs"


class ResumeCustomizer:
    """
    Loads a base resume, customizes it for a specific job using an LLM,
    and generates PDF and DOCX output files.
    """

    def __init__(self):
        RESUME_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        self.llm = None

    def _get_llm(self):
        """Lazy-load LLM client."""
        if self.llm is None:
            from jobflow.utils.llm_client import LLMClient
            self.llm = LLMClient()
        return self.llm

    def load_base_resume(self, path: Optional[str] = None) -> Dict:
        """Load base resume from JSON file."""
        resume_path = Path(path) if path else BASE_RESUME_PATH
        if not resume_path.exists():
            raise FileNotFoundError(f"Base resume not found at {resume_path}")
        with open(resume_path) as f:
            data = json.load(f)
        logger.info(f"Loaded base resume: {data.get('name', 'Unknown')}")
        return data

    def customize_for_job(self, job: Dict, country_config: Dict) -> Dict:
        """
        Use LLM to customize the base resume for a specific job.
        Returns customized resume dict.
        """
        base_resume = self.load_base_resume()
        llm = self._get_llm()

        job_description = job.get("job_description", "")
        job_title = job.get("job_title", "")
        company = job.get("company_name", "")
        country = job.get("country", country_config.get("country", ""))
        required_skills_raw = job.get("required_skills", "[]")
        if isinstance(required_skills_raw, str):
            try:
                required_skills = json.loads(required_skills_raw)
            except Exception:
                required_skills = []
        else:
            required_skills = required_skills_raw

        try:
            customized = llm.customize_resume(
                base_resume=base_resume,
                job_description=f"Job Title: {job_title}\nCompany: {company}\nCountry: {country}\nDescription: {job_description}\nRequired Skills: {', '.join(required_skills)}",
                country=country,
            )

            if isinstance(customized, str):
                try:
                    customized = json.loads(customized)
                except Exception:
                    customized = base_resume.copy()
        except Exception as e:
            logger.error(f"LLM customization failed, using base resume: {e}")
            customized = base_resume.copy()

        customized["_job_id"] = job.get("id", "")
        customized["_customized_at"] = datetime.now().isoformat()
        customized["_customization_for"] = f"{job_title} at {company}"
        return customized

    def generate_pdf(self, resume_data: Dict, output_path: str) -> str:
        """Generate a PDF resume using reportlab."""
        from jobflow.utils.pdf_generator import PDFGenerator
        gen = PDFGenerator()
        gen.generate_resume_pdf(resume_data, output_path)
        logger.info(f"PDF generated: {output_path}")
        return output_path

    def generate_docx(self, resume_data: Dict, output_path: str) -> str:
        """Generate a DOCX resume using python-docx."""
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor, Inches
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = Document()

            # Set margins
            for section in doc.sections:
                section.top_margin = Inches(0.75)
                section.bottom_margin = Inches(0.75)
                section.left_margin = Inches(1.0)
                section.right_margin = Inches(1.0)

            # Name heading
            name_para = doc.add_heading(resume_data.get("name", ""), level=1)
            name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # Contact line
            contact_parts = []
            for field in ("email", "phone", "location", "linkedin"):
                val = resume_data.get(field)
                if val:
                    contact_parts.append(val)
            contact_para = doc.add_paragraph(" | ".join(contact_parts))
            contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

            doc.add_paragraph()

            # Summary
            summary = resume_data.get("summary", "")
            if summary:
                doc.add_heading("Professional Summary", level=2)
                doc.add_paragraph(summary)

            # Skills
            skills = resume_data.get("skills", [])
            if skills:
                doc.add_heading("Skills", level=2)
                skill_text = " • ".join(skills)
                doc.add_paragraph(skill_text)

            # Experience
            experience = resume_data.get("experience", [])
            if experience:
                doc.add_heading("Experience", level=2)
                for exp in experience:
                    p = doc.add_paragraph()
                    run = p.add_run(f"{exp.get('title', '')} — {exp.get('company', '')}")
                    run.bold = True
                    date_p = doc.add_paragraph(
                        f"{exp.get('start_date', '')} - {exp.get('end_date', 'Present')} | {exp.get('location', '')}"
                    )
                    for bullet in exp.get("responsibilities", []):
                        doc.add_paragraph(bullet, style="List Bullet")

            # Education
            education = resume_data.get("education", [])
            if education:
                doc.add_heading("Education", level=2)
                for edu in education:
                    p = doc.add_paragraph()
                    run = p.add_run(f"{edu.get('degree', '')} — {edu.get('institution', '')}")
                    run.bold = True
                    doc.add_paragraph(f"{edu.get('year', '')} | {edu.get('location', '')}")

            # Certifications
            certs = resume_data.get("certifications", [])
            if certs:
                doc.add_heading("Certifications", level=2)
                for cert in certs:
                    doc.add_paragraph(f"• {cert}", style="List Bullet")

            doc.save(output_path)
            logger.info(f"DOCX generated: {output_path}")
            return output_path

        except Exception as e:
            logger.error(f"DOCX generation failed: {e}")
            raise

    def run(self, job_id: str) -> Optional[Dict]:
        """
        Full pipeline: load job, customize resume, generate PDF + DOCX,
        save resume version to database.

        Returns dict with file paths and resume version ID.
        """
        init_db()

        # Load job from database
        job = fetch_one("SELECT * FROM jobs WHERE id = ?", (job_id,))
        if not job:
            logger.error(f"Job not found: {job_id}")
            return None

        # Load country config
        country_code = job.get("country", "UAE").lower()
        config_path = Path(__file__).parent.parent / "configs" / f"country_{country_code}.json"
        country_config = {}
        if config_path.exists():
            with open(config_path) as f:
                country_config = json.load(f)

        # Customize resume
        logger.info(f"Customizing resume for job {job_id}: {job.get('job_title')}")
        customized_resume = self.customize_for_job(job, country_config)

        # Define output paths
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_title = "".join(c if c.isalnum() or c in "-_" else "_" for c in job.get("job_title", "job"))[:30]
        base_name = f"{safe_title}_{timestamp}"
        pdf_path = str(RESUME_OUTPUT_DIR / f"{base_name}.pdf")
        docx_path = str(RESUME_OUTPUT_DIR / f"{base_name}.docx")

        # Generate files
        try:
            self.generate_pdf(customized_resume, pdf_path)
        except Exception as e:
            logger.error(f"PDF generation failed: {e}")
            pdf_path = None

        try:
            self.generate_docx(customized_resume, docx_path)
        except Exception as e:
            logger.error(f"DOCX generation failed: {e}")
            docx_path = None

        # Save to database
        customization_summary = customized_resume.get("_customization_for", "Customized resume")
        version_id = insert_row("resume_versions", {
            "job_id": job_id,
            "version_name": base_name,
            "customization_summary": customization_summary,
            "file_path_pdf": pdf_path,
            "file_path_docx": docx_path,
            "resume_data": json.dumps(customized_resume),
            "country": job.get("country", ""),
        })

        logger.info(f"Resume version {version_id} saved for job {job_id}")
        return {
            "version_id": version_id,
            "pdf_path": pdf_path,
            "docx_path": docx_path,
            "resume_data": customized_resume,
        }


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    customizer = ResumeCustomizer()
    # Test with first available job
    from jobflow.database.db import fetch_all
    jobs = fetch_all("SELECT id, job_title FROM jobs LIMIT 1")
    if jobs:
        result = customizer.run(jobs[0]["id"])
        print(f"Result: {result}")
    else:
        print("No jobs found in database. Run the scraper first.")
